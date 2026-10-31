import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from content import CHEATSHEET, PARTS, REFERENCES, VERIFIED  # noqa: E402

INK = RGBColor(0x14, 0x20, 0x2E)
PAPER = RGBColor(0xEE, 0xF1, 0xF4)
PANEL = RGBColor(0x16, 0x21, 0x2E)
PANEL_TXT = RGBColor(0xE6, 0xEC, 0xF2)
OUT_TXT = RGBColor(0x9F, 0xE0, 0xC9)
SIGNAL = RGBColor(0xD9, 0x48, 0x0F)
AMBER = RGBColor(0xF2, 0xB2, 0x7A)
TEAL = RGBColor(0x0B, 0x6E, 0x61)
MUTED = RGBColor(0x4A, 0x58, 0x68)
DIM = RGBColor(0x8A, 0x9B, 0xAD)
COMMENT = RGBColor(0x7F, 0x91, 0xA4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF6, 0xF8, 0xFA)
RULE = RGBColor(0xC9, 0xD1, 0xDA)
BLANK_FILL = RGBColor(0xFF, 0xF4, 0xE5)
ERR = RGBColor(0xFF, 0xA9, 0x9E)

SANS = "Arial"
MONO = "Courier New"
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)
GROUP_PARTS = range(2, 9)


class Deck:
    def __init__(self, group):
        self.group = group
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.layout = self.prs.slide_layouts[6]
        self.count = 0

    def new(self, bg=PAPER):
        s = self.prs.slides.add_slide(self.layout)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
        self.count += 1
        return s


def box(slide, x, y, w, h, fill=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, paras, size=18, color=INK, font=SANS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, pad=0.08, target=None):
    tb = target or slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(pad))
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        for seg in ([{"t": p}] if isinstance(p, str) else p):
            r = para.add_run()
            r.text = seg["t"]
            f = r.font
            f.name = seg.get("font", font)
            f.size = Pt(seg.get("size", size))
            f.bold = seg.get("bold", bold)
            f.color.rgb = seg.get("color", color)
    return tb


def fit_mono(lines, w_in, h_in, max_pt=18, min_pt=9):
    longest = max((len(line) for line in lines), default=1)
    by_w = (w_in * 72 - 14) / (0.602 * max(longest, 1))
    by_h = (h_in * 72 - 12) / (1.18 * max(len(lines), 1))
    return max(min_pt, min(max_pt, by_w, by_h))


def prompt_lines(code):
    lines = code.split("\n")
    shown = []
    for i, line in enumerate(lines):
        starts = line.strip() and not line.startswith(" ") and not line.startswith("//")
        cont = i > 0 and lines[i - 1].rstrip().endswith((",", "(", "[", "{", ":", ";"))
        shown.append(("> " + line) if starts and not cont else ("  " + line))
    return shown


def code_lines(c):
    return len(c.split("\n"))


def code_panel(slide, x, y, w, h, body, label, color=PANEL_TXT, prompt=True):
    box(slide, x, y, w, h, fill=PANEL)
    text(slide, x, y, w, Inches(0.34), label, size=11, color=DIM, bold=True, pad=0.12)
    shown = prompt_lines(body) if prompt else body.split("\n")
    pt = fit_mono(shown, Emu(w).inches - 0.2, Emu(h).inches - 0.42)
    paras = []
    for line in shown:
        if "<-" in line:
            a, b = line.split("<-", 1)
            paras.append([{"t": a, "color": color}, {"t": "<-" + b, "color": AMBER}])
        elif line.strip().startswith("//"):
            paras.append([{"t": line, "color": COMMENT}])
        elif "Error" in line and not prompt:
            paras.append([{"t": line, "color": ERR}])
        else:
            paras.append([{"t": line if line else " ", "color": color}])
    text(slide, x + Inches(0.06), y + Inches(0.32), w - Inches(0.12), h - Inches(0.36), paras, size=pt, font=MONO, spacing=1.0)


def blank_panel(slide, x, y, w, h):
    shp = box(slide, x, y, w, h, fill=BLANK_FILL, line=SIGNAL)
    text(slide, x, y, w, h, [
        [{"t": "PASTE YOUR SCREENSHOT HERE", "bold": True, "color": SIGNAL, "size": 16}],
        [{"t": "1. Copy the command above into the MongoDB shell, press Enter.", "size": 12}],
        [{"t": "2. Screenshot the result (Windows: Win + Shift + S).", "size": 12}],
        [{"t": "3. Click this box, press Delete, press Ctrl + V, then drag the picture here.", "size": 12}],
    ], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, pad=0.2, target=shp)


def is_bp(deck, part):
    return deck.group and part["n"] in GROUP_PARTS


def header(slide, deck, part, title, soundbite):
    box(slide, 0, 0, Inches(0.18), H, fill=SIGNAL)
    tag = f"PART {part['n']}  ·  {part['title'].upper()}"
    if is_bp(deck, part):
        tag += "  ·  TO DO"
    text(slide, M, Inches(0.32), Inches(11.6), Inches(0.36), tag, size=12, color=SIGNAL, bold=True)
    text(slide, M, Inches(0.62), Inches(12.1), Inches(0.75), title, size=30, bold=True)
    text(slide, M, Inches(1.28), Inches(12.1), Inches(0.5), [[{"t": "Soundbite  ", "bold": True, "color": TEAL}, {"t": soundbite}]], size=16)


def footer(slide, deck, source, verified=False):
    if source:
        text(slide, M, Inches(7.02), Inches(8.0), Inches(0.36), "Source: " + source, size=10, color=MUTED)
    if verified:
        text(slide, Inches(8.4), Inches(7.02), Inches(4.25), Inches(0.36), VERIFIED, size=9, color=TEAL, align=PP_ALIGN.RIGHT)
    text(slide, Inches(12.6), Inches(7.02), Inches(0.5), Inches(0.36), str(deck.count), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def bullets(slide, x, y, w, h, items, size=18):
    text(slide, x, y, w, h, [[{"t": "▸  ", "color": SIGNAL, "bold": True}, {"t": it}] for it in items], size=size, spacing=1.15)


def todo_notes(sl):
    return "TO DO: read the handout page in the footer, then say this in your own words. " + sl["notes"]


def slide_title(deck):
    s = deck.new(bg=INK)
    box(s, 0, Inches(5.9), W, Inches(0.12), fill=SIGNAL)
    text(s, M, Inches(1.2), Inches(12), Inches(0.5), "IT2003 · ADVANCED DATABASE SYSTEMS", size=16, color=AMBER, bold=True)
    text(s, M, Inches(1.8), Inches(12), Inches(2.2), [[{"t": "MongoDB Commands"}], [{"t": "& Cloud Computing"}]], size=60, bold=True, color=WHITE, spacing=0.95)
    sub = ("SHARED GROUP DECK · Part 1 is done. Parts 2-8: find your part and fill the orange boxes."
           if deck.group else "07, 08 and 09 Handouts")
    text(s, M, Inches(4.15), Inches(11.5), Inches(1.2), sub, size=22, color=PANEL_TXT)
    text(s, M, Inches(6.25), Inches(12), Inches(0.9), "Group: ____________________    Section: ________    Date: ____________", size=16, color=PANEL_TXT)
    notes(s, "Introduce the group and the topic: MongoDB commands from the 07 and 08 handouts, and cloud computing from the 09 handout.")


def slide_howto(deck):
    s = deck.new()
    box(s, 0, 0, Inches(0.18), H, fill=SIGNAL)
    text(s, M, Inches(0.4), Inches(12), Inches(0.8), "How to do your part (read this first)", size=34, bold=True)
    steps = [
        ("Find your part", "Slide 3 lists the parts. Type your name next to yours."),
        ("Look at Part 1", "Part 1 is finished. Make yours look like it."),
        ("Open the MongoDB shell", "MongoDB Compass → connect → click  >_ Open MongoDB shell  (top right)."),
        ("Run the setup", "Your part starts with a SETUP slide. Copy all its lines, paste in the shell, press Enter."),
        ("Run each command", "Copy the command on each slide, paste in the shell, press Enter."),
        ("Paste a screenshot", "Win + Shift + S (Mac: Cmd + Shift + 4). Click the orange box → Delete → Ctrl + V → drag the picture into place."),
        ("Fill the blanks", "Replace every ______ with your own words. Use the handout page in the footer."),
        ("Read the speaker notes", "View → Notes. They say what to explain. Practice saying the soundbite."),
    ]
    for i, (head, body) in enumerate(steps):
        cx = M + Inches(6.15) * (i // 4)
        cy = Inches(1.4) + Inches(1.32) * (i % 4)
        box(s, cx, cy, Inches(5.95), Inches(1.18), fill=WHITE, line=RULE)
        text(s, cx + Inches(0.1), cy, Inches(0.8), Inches(1.18), str(i + 1), size=40, bold=True, color=SIGNAL, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + Inches(0.85), cy + Inches(0.08), Inches(5.0), Inches(1.05), [[{"t": head, "bold": True, "size": 18}], [{"t": body, "size": 14, "color": MUTED}]])
    text(s, M, Inches(6.75), Inches(12.1), Inches(0.4), "Only edit YOUR part. Changes save by themselves. Stuck? Message the group leader.", size=15, bold=True, color=SIGNAL)
    footer(s, deck, "")
    notes(s, "Instructions for group members. Delete this slide before presenting.")


def slide_agenda(deck):
    s = deck.new()
    box(s, 0, 0, Inches(0.18), H, fill=SIGNAL)
    text(s, M, Inches(0.5), Inches(12), Inches(0.8), "Eight parts", size=34, bold=True)
    col_w = Inches(5.9)
    for i, p in enumerate(PARTS):
        cx = M + (col_w + Inches(0.3)) * (i // 4)
        cy = Inches(1.55) + Inches(1.32) * (i % 4)
        box(s, cx, cy, col_w, Inches(1.16), fill=WHITE, line=RULE)
        text(s, cx + Inches(0.15), cy + Inches(0.08), Inches(0.8), Inches(1.0), str(p["n"]), size=40, bold=True, color=SIGNAL)
        text(s, cx + Inches(0.95), cy + Inches(0.1), col_w - Inches(1.05), Inches(0.5), p["title"], size=16 if len(p["title"]) > 30 else 19, bold=True)
        text(s, cx + Inches(0.95), cy + Inches(0.55), col_w - Inches(1.05), Inches(0.55), p["kicker"], size=13, color=MUTED, font=MONO)
    footer(s, deck, "07, 08, 09 Handouts")
    notes(s, "Roadmap. Parts 1 to 6 are MongoDB commands from the 07 and 08 handouts. Parts 7 and 8 are cloud computing from the 09 handout.")


def slide_roles(deck):
    s = deck.new()
    box(s, 0, 0, Inches(0.18), H, fill=SIGNAL)
    text(s, M, Inches(0.45), Inches(12), Inches(0.8), "Who presents what", size=34, bold=True)
    tbl = s.shapes.add_table(9, 3, M, Inches(1.45), Inches(12.1), Inches(5.2)).table
    for c, wd in enumerate((Inches(0.9), Inches(7.7), Inches(3.5))):
        tbl.columns[c].width = wd
    for c, hd in enumerate(["Part", "Topic", "Presenter"]):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK
        cell.text = hd
        r = cell.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(16), True, WHITE, SANS
    for i, p in enumerate(PARTS, start=1):
        for c, val in enumerate([str(p["n"]), p["title"] + "  ·  " + p["kicker"], "____________________"]):
            cell = tbl.cell(i, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else ROW_ALT
            cell.text = val
            r = cell.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.name, r.font.color.rgb = Pt(14), SANS, INK
    footer(s, deck, "")
    notes(s, "Type each presenter's name.")


def slide_divider(deck, part):
    s = deck.new(bg=INK)
    text(s, M, Inches(1.4), Inches(4), Inches(2.6), str(part["n"]), size=170, bold=True, color=SIGNAL)
    text(s, Inches(4.2), Inches(0.9), Inches(8.6), Inches(2.0), part["title"], size=40, bold=True, color=WHITE, anchor=MSO_ANCHOR.BOTTOM, spacing=0.95)
    text(s, Inches(4.2), Inches(3.0), Inches(8.6), Inches(0.6), part["kicker"], size=20, color=OUT_TXT, font=MONO)
    text(s, Inches(4.2), Inches(3.8), Inches(8.6), Inches(1.2), "“" + part["soundbite"] + "”", size=28, color=AMBER)
    text(s, Inches(4.2), Inches(5.3), Inches(8.6), Inches(0.6), "Presenter: ____________________", size=20, color=PANEL_TXT)
    n = f"Part {part['n']}: {part['title']}. Soundbite: {part['soundbite']}"
    if is_bp(deck, part):
        n = "TO DO: type your name above. Write a 2-sentence intro for your part here. " + n
    notes(s, n)


def slide_setup(deck, part):
    s = deck.new()
    header(s, deck, part, "SETUP: run this first", "Copy everything in the black box into the MongoDB shell.")
    body = part["setup"]
    code_panel(s, M, Inches(1.95), Inches(12.1), Inches(4.0), body, "COPY ALL OF THIS, PASTE IN THE SHELL, PRESS ENTER", prompt=False)
    bullets(s, M, Inches(6.05), Inches(12.1), Inches(0.9), [
        "This creates the practice data for your part. You can run it again any time to start over.",
        "Delete this slide before presenting.",
    ], size=15)
    footer(s, deck, part["source"])
    notes(s, "Setup only. Delete this slide before presenting.")


def slide_cmd(deck, part, sl, two=False):
    s = deck.new()
    bp = is_bp(deck, part)
    header(s, deck, part, sl["title"], sl["soundbite"])
    top, bottom = Inches(1.9), Inches(6.9)
    if two:
        lw, gap = Inches(6.0), Inches(0.13)
        n = max(code_lines(sl["code"]), code_lines(sl["code2"]))
        ch = Inches(min(2.6, 0.5 + 0.19 * n))
        oh = bottom - top - ch - Inches(0.1)
        for j, (c, o) in enumerate(((sl["code"], sl["output"]), (sl["code2"], sl["output2"]))):
            x = M + (lw + gap) * j
            code_panel(s, x, top, lw, ch, c, "COMMAND")
            if bp:
                blank_panel(s, x, top + ch + Inches(0.1), lw, oh)
            else:
                code_panel(s, x, top + ch + Inches(0.1), lw, oh, o, "OUTPUT", color=OUT_TXT, prompt=False)
    else:
        lw = Inches(7.7)
        ch = Inches(min(2.6, 0.5 + 0.21 * code_lines(sl["code"])))
        code_panel(s, M, top, lw, ch, sl["code"], "COMMAND")
        oy = top + ch + Inches(0.1)
        if bp:
            blank_panel(s, M, oy, lw, bottom - oy)
        else:
            code_panel(s, M, oy, lw, bottom - oy, sl["output"], "OUTPUT", color=OUT_TXT, prompt=False)
        px = M + lw + Inches(0.3)
        pw = W - px - Inches(0.45)
        if bp:
            bullets(s, px, top, pw, Inches(4.9), ["Explain in your own words:"] + ["______________________" for _ in sl["points"]], size=17)
        else:
            bullets(s, px, top, pw, Inches(4.9), sl["points"], size=18)
    footer(s, deck, sl["source"], verified=not bp)
    notes(s, todo_notes(sl) if bp else sl["notes"])


def slide_table(deck, part, sl):
    s = deck.new()
    bp = is_bp(deck, part)
    header(s, deck, part, sl["title"], sl["soundbite"])
    rows, cols = len(sl["rows"]) + 1, len(sl["headers"])
    big = cols <= 3
    tw = W - M - Inches(0.6)
    room = 6.75 - 2.0 - (0.75 if sl.get("points") else 0)
    rh = Inches(min(0.7 if big else 0.95, room / rows))
    tbl = s.shapes.add_table(rows, cols, M, Inches(2.0), tw, rh * rows).table
    for c, hd in enumerate(sl["headers"]):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK
        cell.text = hd
        r = cell.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(18 if big else 16), True, WHITE, SANS
    for r_i, row in enumerate(sl["rows"], start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r_i, c)
            cell.fill.solid()
            blank = bp and c > 0
            cell.fill.fore_color.rgb = BLANK_FILL if blank else (WHITE if r_i % 2 else ROW_ALT)
            cell.text = "________________" if blank else val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(22 if big else 15)
            run.font.name = MONO if (c == 0 and val.startswith("$")) else SANS
            run.font.bold = c == 0
            run.font.color.rgb = SIGNAL if c == 0 and cols > 2 else INK
    if sl.get("points"):
        bullets(s, M, Inches(2.0) + rh * rows + Inches(0.25), tw, Inches(0.8), sl["points"], size=17)
    footer(s, deck, sl["source"])
    notes(s, ("TO DO: fill the orange cells from the handout page in the footer. " + todo_notes(sl)) if bp else sl["notes"])


def slide_text(deck, part, sl):
    s = deck.new()
    bp = is_bp(deck, part)
    header(s, deck, part, sl["title"], sl["soundbite"])
    items = ["______________________________________________" for _ in sl["body"]] if bp else sl["body"]
    bullets(s, M, Inches(2.1), Inches(12), Inches(4.6), items, size=22)
    footer(s, deck, sl["source"])
    notes(s, todo_notes(sl) if bp else sl["notes"])


def slide_benefits(deck, part, sl):
    s = deck.new()
    bp = is_bp(deck, part)
    header(s, deck, part, sl["title"], sl["soundbite"])
    cw = Inches(5.95)
    for i, (name, desc) in enumerate(sl["items"]):
        cx = M + (cw + Inches(0.25)) * (i // 5)
        cy = Inches(2.0) + Inches(0.95) * (i % 5)
        text(s, cx, cy, Inches(0.7), Inches(0.8), str(i + 1), size=30, bold=True, color=SIGNAL, anchor=MSO_ANCHOR.MIDDLE)
        body = ([[{"t": "________________", "bold": True, "size": 18}], [{"t": "one line from 09 Handout p.1-2", "size": 13, "color": MUTED}]]
                if bp else [[{"t": name, "bold": True, "size": 17 if len(name) > 26 else 19}], [{"t": desc, "size": 14, "color": MUTED}]])
        text(s, cx + Inches(0.7), cy, cw - Inches(0.7), Inches(0.8), body)
    footer(s, deck, sl["source"])
    notes(s, ("TO DO: list the 10 benefits in handout order. " + todo_notes(sl)) if bp else sl["notes"])


def slide_cheatsheet(deck):
    s = deck.new()
    box(s, 0, 0, Inches(0.18), H, fill=SIGNAL)
    text(s, M, Inches(0.45), Inches(12), Inches(0.8), "Summary of commands", size=34, bold=True)
    for i, (k, v) in enumerate(CHEATSHEET):
        y = Inches(1.45) + Inches(0.6) * i
        box(s, M, y, Inches(12.1), Inches(0.52), fill=WHITE if i % 2 == 0 else ROW_ALT)
        text(s, M + Inches(0.1), y, Inches(2.0), Inches(0.52), k, size=17, bold=True, color=SIGNAL, anchor=MSO_ANCHOR.MIDDLE)
        text(s, M + Inches(2.1), y, Inches(9.9), Inches(0.52), v, size=15, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, deck, "07, 08, 09 Handouts")
    notes(s, "Leave this up during questions.")


def slide_refs(deck):
    s = deck.new()
    box(s, 0, 0, Inches(0.18), H, fill=INK)
    text(s, M, Inches(0.45), Inches(12), Inches(0.8), "References", size=34, bold=True)
    text(s, M, Inches(1.4), Inches(12.1), Inches(5.4), [[{"t": r}] for r in REFERENCES], size=16, spacing=1.3)
    footer(s, deck, "")
    notes(s, "Sources: the three handouts and the references they cite.")


KINDS = {
    "cmd": lambda d, p, sl: slide_cmd(d, p, sl),
    "cmd2": lambda d, p, sl: slide_cmd(d, p, sl, two=True),
    "table": slide_table,
    "text": slide_text,
    "benefits": slide_benefits,
}


def build(group, out):
    d = Deck(group)
    slide_title(d)
    if group:
        slide_howto(d)
    slide_roles(d)
    slide_agenda(d)
    for part in PARTS:
        slide_divider(d, part)
        if is_bp(d, part) and part.get("setup"):
            slide_setup(d, part)
        for sl in part["slides"]:
            KINDS[sl["kind"]](d, part, sl)
    slide_cheatsheet(d)
    slide_refs(d)
    out.parent.mkdir(parents=True, exist_ok=True)
    d.prs.save(out)
    return d.count


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    n1 = build(False, root / "01_FULL_DECK" / "ADBS_FULL_answer_key.pptx")
    n2 = build(True, root / "02_GROUP_DECK" / "ADBS_GROUP_DECK.pptx")
    print("full", n1, "group", n2)
