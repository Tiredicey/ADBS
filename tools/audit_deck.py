"""Measures a deck: font sizes, words per slide, WCAG contrast of the palette.

  python3 tools/audit_deck.py deck.pptx            new decks: shapes named EXEMPT_* are helper text
  python3 tools/audit_deck.py deck.pptx --legacy   old decks: text below 6.95 in counts as footer
"""
import json
import statistics
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))


def lum(hexcode):
    c = [int(hexcode[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def ratio(a, b):
    hi, lo = sorted((lum(a), lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def runs(shape):
    frames = []
    if shape.has_text_frame:
        frames.append(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    yield r


def main(path, legacy=False, floor=28):
    prs = Presentation(path)
    rows = []
    for i, s in enumerate(prs.slides, 1):
        aud, ex, wc = [], [], 0
        for shp in s.shapes:
            if legacy:
                exempt = shp.top is not None and Emu(shp.top).inches >= 6.95
            else:
                exempt = shp.name.startswith("EXEMPT")
            for r in runs(shp):
                pt = r.font.size.pt if r.font.size else 18.0
                (ex if exempt else aud).append(pt)
                if not exempt:
                    wc += len(r.text.split())
        rows.append({"slide": i, "min_pt": min(aud) if aud else None, "below_floor": sum(1 for x in aud if x < floor),
                     "words": wc, "exempt_min_pt": min(ex) if ex else None})
    words = [r["words"] for r in rows if r["words"]]
    aud_min = [r["min_pt"] for r in rows if r["min_pt"] is not None]
    summary = {
        "deck": Path(path).name, "slides": len(rows), "floor_pt": floor,
        "audience_min_pt": min(aud_min),
        "slides_with_text_below_floor": sum(1 for r in rows if r["below_floor"]),
        "slides_with_text_below_12pt": sum(1 for r in rows if r["min_pt"] is not None and r["min_pt"] < 12),
        "runs_below_floor": sum(r["below_floor"] for r in rows),
        "words_per_slide_median": statistics.median(words), "words_per_slide_max": max(words),
        "exempt_min_pt": min((r["exempt_min_pt"] for r in rows if r["exempt_min_pt"]), default=None),
    }
    out = {"summary": summary, "slides": rows}
    if not legacy:
        from build_deck import CONTRAST_PAIRS, PAL
        pairs = [{"text": a, "bg": b, "use": u, "ratio": round(ratio(PAL[a], PAL[b]), 2)} for a, b, u in CONTRAST_PAIRS]
        out["contrast"] = pairs
        summary["contrast_min"] = min(p["ratio"] for p in pairs)
        summary["contrast_all_pass_AA"] = all(p["ratio"] >= 4.5 for p in pairs)
    return out


if __name__ == "__main__":
    res = main(sys.argv[1], legacy="--legacy" in sys.argv)
    print(json.dumps(res if "--full" in sys.argv else res["summary"], indent=1))
